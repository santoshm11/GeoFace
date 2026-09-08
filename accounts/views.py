import base64
import io
import json
import re
import logging
import os
import traceback

import cv2
import face_recognition
import numpy as np
from PIL import Image
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from django.contrib import messages
from django.contrib.auth import authenticate, get_user_model, login, logout
from django.contrib.auth.decorators import login_required
from django.core.files.base import ContentFile
from django.db import models, transaction
from django.db.models import Q
from django.http import JsonResponse, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.csrf import csrf_exempt

from attendance.models import AttendanceRecord
from attendance.utils import encode_face_from_image
from attendance.utils import get_mandatory_subjects, get_language_subject, get_english_subject

from .forms import (
    AdminUserUpdateForm,
    BulkUploadForm,
    NonTeachingStaffRegistrationForm,
    StudentRegistrationForm,
    TeachingStaffRegistrationForm,
    LecturerAllocationForm,   # new for master login
)
from .models import (
    AcademicYear, Course, Section, StaffProfile,
    StudentEnrollment, StudentProfile, Subject, User
)

User = get_user_model()
logger = logging.getLogger(__name__)

# ---------- Helper: get current academic year ----------
def get_current_academic_year():
    try:
        return AcademicYear.objects.get(is_current=True)
    except AcademicYear.DoesNotExist:
        return None

# ---------- Helper: get expected headers for bulk upload ----------
def get_expected_headers(user_type):
    if user_type == 'student':
        # Added 'academic_year' (optional) – we can default to current
        return ['full_name', 'uucms_id', 'email', 'phone_number',
                'course_code', 'semester', 'section_name', 'language1', 'academic_year']
    elif user_type == 'lecturer':
        return ['full_name', 'username', 'email', 'phone_number', 'designation', 'course_code']
    elif user_type == 'staff':
        return ['full_name', 'username', 'email', 'phone_number', 'designation']
    return []

# ---------- Bulk row processing (updated) ----------
def process_bulk_row(user_type, row, results):
    from accounts.models import User, StudentProfile, Course, Section, StaffProfile, StudentEnrollment, Subject, AcademicYear
    from attendance.utils import get_mandatory_subjects, get_language_subject, get_english_subject
    from django.db import transaction
    import logging
    logger = logging.getLogger(__name__)

    def clean_text(val):
        return str(val).strip() if val else ''

    if user_type == 'student':
        try:
            # Expect 9 columns now (academic_year optional)
            if len(row) < 8:
                raise ValueError(f"Insufficient columns. Expected at least 8, got {len(row)}")
            full_name, uucms_id, email, phone, course_code, semester, section_name, language1 = row[:8]
            academic_year_name = clean_text(row[8]) if len(row) > 8 else ''

            full_name = clean_text(full_name)
            uucms_id = clean_text(uucms_id)
            email = clean_text(email)
            phone = clean_text(phone)
            course_code = clean_text(course_code)
            semester = int(clean_text(semester)) if clean_text(semester) else None
            section_name = clean_text(section_name)
            language1 = clean_text(language1).upper()

            if not all([full_name, uucms_id, email, course_code, semester, section_name, language1]):
                raise ValueError("Missing required fields")

            # Normalize language1
            if language1 in ['KAN', 'KANNADA']:
                language1 = 'KAN'
            elif language1 in ['HIN', 'HINDI']:
                language1 = 'HIN'
            else:
                raise ValueError("Language1 must be KAN or HIN (or Kannada/Hindi)")

            if User.objects.filter(username=uucms_id).exists():
                raise ValueError(f"UUCMS ID '{uucms_id}' already exists")
            if StudentProfile.objects.filter(roll_no=uucms_id).exists():
                raise ValueError(f"Roll number '{uucms_id}' already exists")

            # Find course
            course = Course.objects.filter(code=course_code).first()
            if not course:
                course = Course.objects.filter(name__iexact=course_code).first()
            if not course and '-' in course_code:
                parts = course_code.split('-')
                if parts[0].strip():
                    course = Course.objects.filter(code=parts[0].strip()).first()
            if not course:
                raise ValueError(f"Course '{course_code}' not found")

            section = Section.objects.filter(course=course, semester=semester, name=section_name).first()
            if not section:
                raise ValueError(f"Section '{section_name}' not found for course {course.code} semester {semester}")

            # Determine academic year
            academic_year = None
            if academic_year_name:
                academic_year = AcademicYear.objects.filter(name__iexact=academic_year_name).first()
            if not academic_year:
                academic_year = get_current_academic_year()
            if not academic_year:
                raise ValueError("No academic year provided and no current year set.")

            with transaction.atomic():
                user = User.objects.create_user(
                    username=uucms_id,
                    password='student123',
                    first_name=full_name,
                    last_name='',
                    email=email,
                    phone_number=phone,
                    user_type='student'
                )
                profile = StudentProfile.objects.create(
                    user=user,
                    roll_no=uucms_id,
                    department=course.department,
                    registration_complete=False
                )
                enrollment = StudentEnrollment.objects.create(
                    student=profile,
                    course=course,
                    academic_year=academic_year,
                    semester=semester,
                    section=section,
                    language1=language1
                )
                # Assign subjects
                core = get_mandatory_subjects(course, semester, section)
                lang1 = get_language_subject(course, semester, language1)
                lang2 = get_english_subject(course, semester)
                all_subjects = list(core) + [lang1, lang2]
                enrollment.subjects.add(*all_subjects)
                enrollment.save()

            results['success'] += 1

        except Exception as e:
            results['failed'] += 1
            results['errors'].append(f"Student row: {str(e)}")
            raise

    elif user_type == 'lecturer':
        # Same as before, but we might add academic_year if needed – not required for lecturers
        try:
            if len(row) < 6:
                raise ValueError(f"Insufficient columns. Expected 6, got {len(row)}")
            full_name, username, email, phone, designation, course_code = row[:6]
            full_name = clean_text(full_name)
            username = clean_text(username)
            email = clean_text(email)
            phone = clean_text(phone)
            designation = clean_text(designation)
            course_code = clean_text(course_code)

            if not all([full_name, username, email, designation]):
                raise ValueError("Missing required fields")
            if User.objects.filter(username=username).exists():
                raise ValueError(f"Username '{username}' already exists")

            course = None
            if course_code:
                course = Course.objects.filter(code=course_code).first()
                if not course:
                    raise ValueError(f"Assigned course '{course_code}' not found")

            with transaction.atomic():
                user = User.objects.create_user(
                    username=username,
                    password='staff123',
                    first_name=full_name,
                    last_name='',
                    email=email,
                    phone_number=phone,
                    user_type='lecturer',
                    assigned_course=course
                )
                StaffProfile.objects.create(
                    user=user,
                    employee_id=username,
                    department=course.name if course else '',
                    designation=designation,
                    registration_complete=False
                )
            results['success'] += 1
        except Exception as e:
            results['failed'] += 1
            results['errors'].append(f"Lecturer row: {str(e)}")
            raise

    elif user_type == 'staff':
        try:
            if len(row) < 5:
                raise ValueError(f"Insufficient columns. Expected 5, got {len(row)}")
            full_name, username, email, phone, designation = row[:5]
            full_name = clean_text(full_name)
            username = clean_text(username)
            email = clean_text(email)
            phone = clean_text(phone)
            designation = clean_text(designation)

            if not all([full_name, username, email, designation]):
                raise ValueError("Missing required fields")
            if User.objects.filter(username=username).exists():
                raise ValueError(f"Username '{username}' already exists")

            with transaction.atomic():
                user = User.objects.create_user(
                    username=username,
                    password='staff123',
                    first_name=full_name,
                    last_name='',
                    email=email,
                    phone_number=phone,
                    user_type='staff'
                )
                StaffProfile.objects.create(
                    user=user,
                    employee_id=username,
                    department='',
                    designation=designation,
                    registration_complete=False
                )
            results['success'] += 1
        except Exception as e:
            results['failed'] += 1
            results['errors'].append(f"Staff row: {str(e)}")
            raise

    else:
        raise ValueError(f"Unsupported user type: {user_type}")

# ---------- Redirect helper ----------
def redirect_to_dashboard(request, user):
    dashboard_map = {
        'principal': 'principal_dashboard',
        'hod': 'hod_dashboard',
        'lecturer': 'lecturer_dashboard',
        'staff': 'staff_dashboard',
        'admin': 'admin_dashboard',
    }
    if user.user_type in dashboard_map:
        return redirect(dashboard_map[user.user_type])
    if user.user_type == 'student':
        return redirect('student_dashboard')
    return redirect('home')

# ---------- Authentication views ----------
def login_view(request):
    if request.user.is_authenticated:
        return redirect_to_dashboard(request, request.user)
    if request.method == 'POST':
        username = request.POST.get('username', '').strip()
        password = request.POST.get('password', '')
        try:
            user_obj = User.objects.get(username__iexact=username)
            actual_username = user_obj.username
        except User.DoesNotExist:
            actual_username = username
        user = authenticate(request, username=actual_username, password=password)
        if user is not None:
            login(request, user)
            messages.success(request, f'Welcome back, {user.get_full_name() or user.username}!')
            return redirect_to_dashboard(request, user)
        else:
            messages.error(request, 'Invalid username or password.')
    return render(request, 'accounts/login.html')

def logout_view(request):
    logout(request)
    messages.success(request, 'You have been logged out successfully.')
    return redirect('login')

def register_view(request):
    # Keep as is – uses old self-registration (not used in new schema)
    # We'll keep it for backward compatibility
    if request.user.is_authenticated:
        return redirect('home')
    if request.method == 'POST':
        # ... unchanged (original code)
        pass
    return render(request, 'accounts/register.html')

# ---------- Face registration and profile views ----------
def handle_face_image(profile, face_image):
    """Update face image and encoding for a profile."""
    try:
        face_image.seek(0)
        if profile.face_image:
            profile.face_image.delete(save=False)
        profile.face_image = face_image
        profile.save(update_fields=['face_image'])
        encoding = encode_face_from_image(profile.face_image.path)
        if encoding is not None:
            profile.set_face_encoding(encoding)
            profile.registration_complete = True
            profile.save(update_fields=['face_encoding', 'registration_complete'])
            return True, "Face image and encoding updated successfully."
        else:
            profile.face_image.delete(save=False)
            profile.face_encoding = None
            profile.registration_complete = False
            profile.save(update_fields=['face_encoding', 'registration_complete'])
            return False, "No face detected in the uploaded image. Please try again."
    except Exception as e:
        logger.exception("Error in handle_face_image")
        return False, f"Error processing face image: {str(e)}"

@login_required
def face_registration_view(request):
    if request.user.user_type != 'student':
        messages.error(request, 'Only students can register face.')
        return redirect('home')
    student_profile = get_object_or_404(StudentProfile, user=request.user)
    if request.method == 'POST':
        face_image = request.FILES.get('face_image')
        if face_image:
            success, msg = handle_face_image(student_profile, face_image)
            if success:
                messages.success(request, msg)
                return redirect('student_dashboard')
            else:
                messages.error(request, msg)
        else:
            face_data = request.POST.get('face_data')
            if face_data:
                try:
                    if ',' in face_data:
                        face_data = face_data.split(',')[1]
                    img_bytes = base64.b64decode(face_data)
                    np_arr = np.frombuffer(img_bytes, np.uint8)
                    img = cv2.imdecode(np_arr, cv2.IMREAD_COLOR)
                    rgb_img = cv2.cvtColor(img, cv2.COLOR_BGR2RGB)
                    pil_img = Image.fromarray(rgb_img)
                    img_io = io.BytesIO()
                    pil_img.save(img_io, format='JPEG')
                    img_io.seek(0)
                    success, msg = handle_face_image(student_profile, img_io)
                    if success:
                        return JsonResponse({'success': True, 'message': msg})
                    else:
                        return JsonResponse({'success': False, 'message': msg})
                except Exception as e:
                    return JsonResponse({'success': False, 'message': f'Error: {str(e)}'})
    return render(request, 'accounts/face_registration.html', {'student_profile': student_profile})

@login_required
@csrf_exempt
def update_face_encoding(request):
    if request.user.user_type != 'student':
        return JsonResponse({'error': 'Unauthorized'}, status=403)
    if request.method != 'POST':
        return JsonResponse({'error': 'Invalid method'}, status=405)
    try:
        data = json.loads(request.body)
        face_data = data.get('face_data')
        if not face_data:
            return JsonResponse({'success': False, 'message': 'No face data provided.'})
        if ',' in face_data:
            face_data = face_data.split(',')[1]
        img_bytes = base64.b64decode(face_data)
        np_arr = np.frombuffer(img_bytes, np.uint8)
        img = cv2.imdecode(np_arr, cv2.IMREAD_COLOR)
        rgb_img = cv2.cvtColor(img, cv2.COLOR_BGR2RGB)
        face_encodings = face_recognition.face_encodings(rgb_img)
        if not face_encodings:
            return JsonResponse({'success': False, 'message': 'No face detected.'})
        student_profile = get_object_or_404(StudentProfile, user=request.user)
        student_profile.set_face_encoding(face_encodings[0])
        student_profile.save()
        return JsonResponse({'success': True, 'message': 'Face encoding updated.'})
    except Exception as e:
        return JsonResponse({'success': False, 'message': f'Error: {str(e)}'})

@login_required
def profile_view(request):
    context = {}
    if request.user.user_type == 'student':
        try:
            context['student_profile'] = StudentProfile.objects.get(user=request.user)
        except StudentProfile.DoesNotExist:
            pass
    return render(request, 'accounts/profile.html', context)

@login_required
def edit_profile(request):
    if request.method == 'POST':
        user = request.user
        user.first_name = request.POST.get('first_name', user.first_name).strip()
        user.last_name = request.POST.get('last_name', user.last_name).strip()
        user.email = request.POST.get('email', user.email).strip()
        user.phone_number = request.POST.get('phone_number', user.phone_number).strip()
        user.save()
        if user.user_type == 'student':
            try:
                student_profile = StudentProfile.objects.get(user=user)
                student_profile.department = request.POST.get('department', student_profile.department).strip()
                student_profile.save()
            except StudentProfile.DoesNotExist:
                pass
        messages.success(request, 'Profile updated successfully!')
        return redirect('profile')
    context = {}
    if request.user.user_type == 'student':
        try:
            context['student_profile'] = StudentProfile.objects.get(user=request.user)
        except StudentProfile.DoesNotExist:
            pass
    return render(request, 'accounts/edit_profile.html', context)

@login_required
def change_password(request):
    if request.method == 'POST':
        user = request.user
        old_password = request.POST.get('old_password')
        new_password = request.POST.get('new_password')
        confirm_password = request.POST.get('confirm_password')
        if not user.check_password(old_password):
            messages.error(request, 'Current password is incorrect.')
            return redirect('change_password')
        if new_password != confirm_password:
            messages.error(request, 'New passwords do not match.')
            return redirect('change_password')
        if len(new_password) < 8:
            messages.error(request, 'Password must be at least 8 characters long.')
            return redirect('change_password')
        user.set_password(new_password)
        user.save()
        login(request, user)
        messages.success(request, 'Password changed successfully!')
        return redirect('profile')
    return render(request, 'accounts/change_password.html')

# ---------- Lecturer views (unchanged) ----------
@login_required
def student_list(request):
    if request.user.user_type != 'lecturer':
        messages.error(request, 'Access denied.')
        return redirect('home')
    students = StudentProfile.objects.select_related('user').all()
    return render(request, 'accounts/student_list.html', {'students': students})

@login_required
def student_detail(request, student_id):
    if request.user.user_type != 'lecturer':
        messages.error(request, 'Access denied.')
        return redirect('home')
    student = get_object_or_404(StudentProfile.objects.select_related('user'), id=student_id)
    total_classes = AttendanceRecord.objects.filter(student=student).count()
    present_count = AttendanceRecord.objects.filter(student=student, status='Present').count()
    attendance_percentage = (present_count / total_classes * 100) if total_classes > 0 else 0
    context = {
        'student': student,
        'total_classes': total_classes,
        'present_count': present_count,
        'attendance_percentage': round(attendance_percentage, 2),
    }
    return render(request, 'accounts/student_detail.html', context)

# ---------- Admin views ----------
@login_required
def admin_dashboard(request):
    if request.user.user_type != 'admin':
        messages.error(request, "Access denied.")
        return redirect('dashboard')
    context = {
        'total_students': User.objects.filter(user_type='student').count(),
        'total_teaching': User.objects.filter(user_type='lecturer').count(),
        'total_nonteaching': User.objects.filter(user_type='staff').count(),
    }
    return render(request, 'admin/admin_dashboard.html', context)

@login_required
def admin_register_choice(request):
    if request.user.user_type != 'admin':
        messages.error(request, "Access denied.")
        return redirect('admin_dashboard')
    return render(request, 'admin/admin_register_choice.html')

# ----- Student Registration (using updated form) -----
@login_required
def admin_register_student(request):
    if request.user.user_type not in ['admin', 'principal']:
        messages.error(request, 'Access denied.')
        return redirect('home')

    if request.method == 'POST' and request.POST.get('confirm') == '1':
        temp_data = request.session.get('temp_student_data', {})
        if not temp_data:
            messages.error(request, 'Session expired. Please start over.')
            return redirect('admin_register_student')

        # Recreate form from session data
        form = StudentRegistrationForm(temp_data)
        if form.is_valid():
            try:
                user = form.save()
                messages.success(request, f'Student {user.get_full_name()} registered successfully!')
                # Clear session
                request.session.pop('temp_student_data', None)
                request.session.pop('temp_face_image', None)
                request.session.pop('temp_face_image_name', None)
                return redirect('admin_dashboard')
            except Exception as e:
                messages.error(request, f'Error saving student: {str(e)}')
                return redirect('admin_register_student')
        else:
            messages.error(request, 'Invalid data. Please try again.')
            return redirect('admin_register_student')

    # First POST (show confirmation)
    if request.method == 'POST':
        form = StudentRegistrationForm(request.POST, request.FILES)
        if form.is_valid():
            data = form.cleaned_data
            # Store in session for confirmation
            request.session['temp_student_data'] = {
                'student_name': data['student_name'],
                'uucms_id': data['uucms_id'],
                'email': data['email'],
                'phone_number': data.get('phone_number', ''),
                'password': data['password'],
                'course': data['course'].id,
                'semester': data['semester'],
                'section': data['section'].id,
                'language1': data['language1'],
                'academic_year': data['academic_year'].id if data.get('academic_year') else None,
            }

            # Store face image if uploaded
            if request.FILES.get('face_image'):
                img = request.FILES['face_image']
                img.seek(0)
                encoded = base64.b64encode(img.read()).decode('utf-8')
                request.session['temp_face_image'] = encoded
                request.session['temp_face_image_name'] = img.name

            # Prepare confirmation data
            core_subjects = get_mandatory_subjects(data['course'], data['semester'], data['section'])
            lang_subject = get_language_subject(data['course'], data['semester'], data['language1'])
            english_subject = get_english_subject(data['course'], data['semester'])
            all_subjects = list(core_subjects) + [lang_subject, english_subject]
            subject_names = [f"{s.code} - {s.name}" for s in all_subjects]

            confirm_data = {
                'full_name': data['student_name'],
                'uucms_id': data['uucms_id'],
                'email': data['email'],
                'phone': data.get('phone_number', ''),
                'course': data['course'],
                'semester': data['semester'],
                'section': data['section'],
                'language1': 'Kannada' if data['language1'] == 'KAN' else 'Hindi',
                'language2': 'Basic English',
                'subjects': subject_names,
                'academic_year': data['academic_year'].name if data.get('academic_year') else 'Current',
            }
            context = {
                'form': form,
                'confirm_data': confirm_data,
                'show_confirm': True,
            }
            return render(request, 'admin/register_student.html', context)
        else:
            context = {'form': form, 'show_confirm': False}
            return render(request, 'admin/register_student.html', context)

    # GET
    form = StudentRegistrationForm()
    context = {'form': form, 'show_confirm': False}
    return render(request, 'admin/register_student.html', context)

# ----- Teaching Staff Registration (using updated form) -----
@login_required
def admin_register_teaching(request):
    if request.user.user_type != 'admin':
        messages.error(request, "Access denied.")
        return redirect('admin_dashboard')

    if request.method == 'POST':
        form = TeachingStaffRegistrationForm(request.POST, request.FILES)
        if form.is_valid():
            try:
                user = form.save()
                messages.success(request, f"Teaching staff {user.get_full_name()} registered.")
                return redirect('admin_teaching_list')
            except Exception as e:
                messages.error(request, f"Error: {str(e)}")
                return render(request, 'admin/admin_register_teaching.html', {'form': form})
        else:
            return render(request, 'admin/admin_register_teaching.html', {'form': form})
    else:
        form = TeachingStaffRegistrationForm()
        return render(request, 'admin/admin_register_teaching.html', {'form': form})

# ----- Non-Teaching Staff Registration (using updated form) -----
@login_required
def admin_register_nonteaching(request):
    if request.user.user_type != 'admin':
        messages.error(request, "Access denied.")
        return redirect('admin_dashboard')

    if request.method == 'POST':
        form = NonTeachingStaffRegistrationForm(request.POST, request.FILES)
        if form.is_valid():
            try:
                user = form.save()
                messages.success(request, f"Non-teaching staff {user.get_full_name()} registered.")
                return redirect('admin_nonteaching_list')
            except Exception as e:
                messages.error(request, f"Error: {str(e)}")
                return render(request, 'admin/admin_register_nonteaching.html', {'form': form})
        else:
            return render(request, 'admin/admin_register_nonteaching.html', {'form': form})
    else:
        form = NonTeachingStaffRegistrationForm()
        return render(request, 'admin/admin_register_nonteaching.html', {'form': form})

# ----- Student List (Admin) -----
@login_required
def admin_student_list(request):
    if request.user.user_type != 'admin':
        messages.error(request, "Access denied.")
        return redirect('admin_dashboard')

    enrollments = StudentEnrollment.objects.select_related(
        'student__user', 'course', 'section', 'academic_year'
    ).all()

    search = request.GET.get('search', '')
    course_id = request.GET.get('course')
    year = request.GET.get('year')
    section_name = request.GET.get('section')
    academic_year_id = request.GET.get('academic_year')

    if search:
        enrollments = enrollments.filter(
            Q(student__user__first_name__icontains=search) |
            Q(student__user__last_name__icontains=search) |
            Q(student__user__username__icontains=search) |
            Q(student__roll_no__icontains=search)
        )
    if course_id:
        enrollments = enrollments.filter(course_id=course_id)
    if year:
        enrollments = enrollments.filter(semester=year)
    if section_name:
        enrollments = enrollments.filter(section__name=section_name)
    if academic_year_id:
        enrollments = enrollments.filter(academic_year_id=academic_year_id)

    courses = Course.objects.all()
    sections = Section.objects.values_list('name', flat=True).distinct().order_by('name')
    academic_years = AcademicYear.objects.all().order_by('-start_date')

    context = {
        'enrollments': enrollments,
        'courses': courses,
        'sections': sections,
        'academic_years': academic_years,
        'search': search,
        'selected_course': course_id,
        'selected_year': year,
        'selected_section': section_name,
        'selected_academic_year': academic_year_id,
    }
    return render(request, 'admin/admin_student_list.html', context)

# ----- User Detail (Admin) -----
@login_required
def admin_user_detail(request, user_id):
    if request.user.user_type != 'admin':
        messages.error(request, "Access denied.")
        return redirect('admin_dashboard')

    user_obj = get_object_or_404(User, id=user_id)
    profile = None
    if user_obj.user_type == 'student':
        try:
            profile = user_obj.student_profile
        except StudentProfile.DoesNotExist:
            pass

    if request.method == 'POST':
        if 'delete' in request.POST:
            user_obj.delete()
            messages.success(request, "User deleted.")
            return redirect('admin_dashboard')

        form = AdminUserUpdateForm(request.POST, instance=user_obj)
        if form.is_valid():
            form.save()
            messages.success(request, "User updated.")

            face_image = request.FILES.get('face_image')
            if face_image:
                if user_obj.user_type == 'student' and profile:
                    success, msg = handle_face_image(profile, face_image)
                elif user_obj.user_type in ['lecturer', 'hod', 'staff']:
                    staff_profile = getattr(user_obj, 'staff_profile', None)
                    if staff_profile:
                        success, msg = handle_face_image(staff_profile, face_image)
                    else:
                        success, msg = False, "Staff profile not found."
                else:
                    success, msg = False, "Face image not applicable for this user type."
                if success:
                    messages.success(request, msg)
                else:
                    messages.warning(request, f"Face update issue: {msg}")
            return redirect('admin_user_detail', user_id=user_obj.id)
        else:
            messages.error(request, "Please correct the errors below.")
    else:
        form = AdminUserUpdateForm(instance=user_obj)

    # Get enrollments for student
    enrollments = None
    if user_obj.user_type == 'student' and profile:
        enrollments = StudentEnrollment.objects.filter(student=profile).select_related('course', 'section', 'academic_year')

    context = {
        'user_obj': user_obj,
        'profile': profile,
        'form': form,
        'enrollments': enrollments,
    }
    return render(request, 'admin/admin_user_detail.html', context)

# ----- Teaching Staff List -----
@login_required
def admin_teaching_list(request):
    if request.user.user_type != 'admin':
        messages.error(request, "Access denied.")
        return redirect('admin_dashboard')

    staff = User.objects.filter(user_type='lecturer')
    search = request.GET.get('search', '')
    course_id = request.GET.get('course')
    department = request.GET.get('department')
    designation = request.GET.get('designation')

    if search:
        staff = staff.filter(
            Q(first_name__icontains=search) |
            Q(last_name__icontains=search) |
            Q(username__icontains=search) |
            Q(email__icontains=search)
        )
    if course_id:
        staff = staff.filter(assigned_course_id=course_id)
    if department:
        staff = staff.filter(staff_profile__department__icontains=department)
    if designation:
        staff = staff.filter(staff_profile__designation=designation)

    courses = Course.objects.all()
    departments = StaffProfile.objects.filter(
        user__user_type='lecturer'
    ).values_list('department', flat=True).distinct().exclude(department='')

    TEACHING_DESIGNATIONS = ['Professor', 'Associate Professor', 'Assistant Professor', 'Lecturer/Instructor']
    existing = StaffProfile.objects.filter(
        user__user_type='lecturer'
    ).values_list('designation', flat=True).distinct().exclude(designation='')
    designations = TEACHING_DESIGNATIONS.copy()
    for d in existing:
        if d not in designations:
            designations.append(d)

    context = {
        'staff': staff,
        'search': search,
        'courses': courses,
        'selected_course': course_id,
        'departments': departments,
        'selected_department': department,
        'designations': designations,
        'selected_designation': designation,
    }
    return render(request, 'admin/admin_teaching_staff_list.html', context)

# ----- Non-Teaching Staff List -----
@login_required
def admin_nonteaching_list(request):
    if request.user.user_type != 'admin':
        messages.error(request, "Access denied.")
        return redirect('admin_dashboard')

    staff = User.objects.filter(user_type='staff')
    search = request.GET.get('search', '')
    designation = request.GET.get('designation')

    if search:
        staff = staff.filter(
            Q(first_name__icontains=search) |
            Q(last_name__icontains=search) |
            Q(username__icontains=search) |
            Q(email__icontains=search)
        )
    if designation:
        staff = staff.filter(staff_profile__designation=designation)

    NON_TEACHING_DESIGNATIONS = ['FDA', 'SDA', 'Peon', 'Lab Attender', 'Sweeper']

    context = {
        'staff': staff,
        'search': search,
        'type': 'Non-Teaching',
        'designations': NON_TEACHING_DESIGNATIONS,
        'selected_designation': designation,
    }
    return render(request, 'admin/admin_nonteaching_staff_list.html', context)

# ----- Bulk Registration -----
@login_required
def admin_bulk_register(request):
    if request.user.user_type != 'admin':
        messages.error(request, "Access denied.")
        return redirect('admin_dashboard')

    if request.method == 'POST':
        form = BulkUploadForm(request.POST, request.FILES)
        if form.is_valid():
            user_type = form.cleaned_data['user_type']
            excel_file = request.FILES['excel_file']
            logger.info(f"BULK REGISTRATION START: user_type={user_type}, file={excel_file.name}")

            try:
                if not excel_file.name.endswith(('.xlsx', '.xls')):
                    raise ValueError("File must be an Excel file (.xlsx or .xls).")

                wb = load_workbook(excel_file)
                ws = wb.active
                rows = list(ws.iter_rows(values_only=True))
                if len(rows) < 2:
                    raise ValueError("File is empty or missing data rows.")

                headers = rows[0]
                data = rows[1:]
                expected = get_expected_headers(user_type)
                headers_clean = [str(h).strip() for h in headers]
                expected_clean = [e.strip() for e in expected]

                if headers_clean != expected_clean:
                    raise ValueError(f"Invalid headers. Expected: {', '.join(expected)}. Got: {', '.join(headers_clean)}")

                results = {'success': 0, 'failed': 0, 'errors': []}
                for row_idx, row in enumerate(data, start=2):
                    if not any(cell is not None and str(cell).strip() for cell in row):
                        continue
                    try:
                        process_bulk_row(user_type, row, results)
                    except Exception as e:
                        results['failed'] += 1
                        results['errors'].append(f"Row {row_idx}: {str(e)}")
                        logger.error(f"Row {row_idx} error: {e}")

                logger.info(f"Bulk results: Success={results['success']}, Failed={results['failed']}")
                messages.success(request, f"Bulk registration completed. Success: {results['success']}, Failed: {results['failed']}.")
                if results['errors']:
                    errors_str = '; '.join(results['errors'][:10])
                    if len(results['errors']) > 10:
                        errors_str += f" and {len(results['errors'])-10} more."
                    messages.warning(request, f"Errors: {errors_str}")
                return redirect('admin_dashboard')

            except Exception as e:
                logger.exception("Bulk upload error")
                messages.error(request, f"Error processing file: {str(e)}")
                return render(request, 'admin/admin_bulk_register.html', {'form': form})
        else:
            messages.error(request, "Invalid form data.")
    else:
        form = BulkUploadForm()

    return render(request, 'admin/admin_bulk_register.html', {'form': form})

# ----- Download Sample Excel -----
@login_required
def admin_download_sample(request, user_type):
    if request.user.user_type != 'admin':
        messages.error(request, "Access denied.")
        return redirect('admin_dashboard')

    headers = get_expected_headers(user_type)
    if not headers:
        return HttpResponse("Invalid user type", status=400)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sample Data"

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = min(len(header) + 4, 30)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename=sample_{user_type}.xlsx'
    wb.save(response)
    return response